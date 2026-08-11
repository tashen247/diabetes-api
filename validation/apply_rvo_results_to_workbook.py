#!/usr/bin/env python3
"""Apply C:\\AI_Images\\RVO\\eval_rvo_results.json into Fundus_AI_Validation.xlsx."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", required=True, help="Path to eval_rvo_results.json")
    ap.add_argument(
        "--workbook",
        default=str(Path(__file__).with_name("Fundus_AI_Validation.xlsx")),
    )
    args = ap.parse_args()

    payload = json.loads(Path(args.json).read_text(encoding="utf-8"))
    val = payload["val"]
    split = payload["split"]
    wb = openpyxl.load_workbook(args.workbook)

    ws = wb["Summary"]
    ws["H8"] = int(split["n_train"])
    ws["I8"] = int(split["n_pos_train"])
    ws["J8"] = "post-hoc val 10% (random_state=42)"
    ws["K8"] = int(split["n_val"])
    ws["L8"] = int(split["n_pos_val"])
    ws["M8"] = round(float(val["sensitivity"]), 4)
    ws["N8"] = round(float(val["specificity"]), 4)
    ws["O8"] = round(float(val["auroc"]), 4)
    ws["P8"] = (
        f"Train AUROC {payload['train']['auroc']:.4f}. "
        f"Argmax@0.5. Quote val AUROC. Root={payload['data_root']}"
    )
    ws["Q8"] = "AUROC on post-hoc 10% of C:\\AI_Images\\RVO (GPU eval_rvo_ai_images.py)"
    ws["R8"] = (
        f"pos_dir={payload['pos_dir']}; neg_dir={payload['neg_dir']}; "
        f"n_labelled={payload['n_labelled']}. Quote val {val['auroc']:.4f}, not max(train,val)."
    )
    ws["E8"] = f"{payload['data_root']} ({payload['layout']})"
    ws["F8"] = payload.get("assembled_utc", "")[:10]

    ws = wb["RVO"]
    ws["B6"] = int(payload["n_pos_total"])
    ws["C6"] = int(payload["n_neg_total"])
    ws["D6"] = payload.get("assembled_utc", "")[:10]
    ws["E6"] = "post-hoc 10%"
    ws["B11"] = round(float(val["sensitivity"]), 4)
    ws["C11"] = "argmax @0.5 on val"
    ws["B12"] = round(float(val["specificity"]), 4)
    ws["C12"] = "argmax @0.5 on val"
    ws["B13"] = round(float(val["auroc"]), 4)
    ws["C13"] = f"val n={split['n_val']} n_pos={split['n_pos_val']}"
    ws["B14"] = "post-hoc 10% (random_state=42)"
    ws["C14"] = "filled from eval JSON"

    ws = wb["Holdout_eval"]
    ws["B17"] = "post-hoc val 10%"
    ws["C17"] = int(split["n_val"])
    ws["D17"] = int(split["n_pos_val"])
    ws["E17"] = round(float(val["auroc"]), 4)
    ws["F17"] = round(float(val["sensitivity"]), 4)
    ws["G17"] = round(float(val["specificity"]), 4)
    ws["H17"] = f"from {payload['data_root']}; quote AUROC"

    ws = wb["Gaps"]
    ws["B8"] = "None (AUROC filled from C:\\AI_Images\\RVO)"
    ws["C8"] = "Optional: Youden threshold / independent clinical hold-out"
    ws["D8"] = f"Filled val AUROC {val['auroc']:.4f}"

    if "Finish_status" in wb.sheetnames:
        ws = wb["Finish_status"]
        ws["B8"] = f"{val['auroc']:.4f} (post-hoc val)"
        ws["C8"] = "filled (argmax)"
        ws["D8"] = "DONE — quote val AUROC from C:\\AI_Images\\RVO"

    wb.save(args.workbook)
    print(f"Updated {args.workbook} with RVO val AUROC={val['auroc']:.4f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
